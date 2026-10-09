"""Reusable ME2M processor based on the supplied ME2M/build_report.py.

Reading, classification, summary and export are independent of the ACR
pipeline. Importing this module never reads or writes a workbook.
"""
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CATEGORIES = ("Discontinued", "Past Dues", "Product in McCook", "With NCR", "(blank)")
REQUIRED_COLUMNS = (
    "Purchasing Group", "Delivery Date", "Delivery Text", "ABC Indicator", "Net Order Value",
)
NUMBER_FORMATS = {
    "Material": "0", "Purchasing Document": "0", "Item": "#,##0",
    "Plant": "#,##0", "Purchasing Organization": "#,##0", "Purch. Organization": "#,##0",
    "Order Quantity": "#,##0.00", "Still to be delivered (qty)": "#,##0.00",
    "Still to be invoiced (val.)": "#,##0.00", "Still to be delivered (value)": "#,##0.00",
    "Net Order Value": "#,##0.00",
}
MEASURE_COLUMNS = {name for name, fmt in NUMBER_FORMATS.items() if fmt.endswith(".00")}
REPORT_TIMEZONE = timezone(timedelta(hours=5, minutes=30))


class ME2MCancelled(Exception):
    """Raised at a stage boundary when the user stops a report."""


def report_today():
    return datetime.now(REPORT_TIMEZONE).date()


def _read_source(source):
    with pd.ExcelFile(source, engine="openpyxl") as book:
        if not book.sheet_names:
            raise ValueError("The input workbook must contain a worksheet.")
        # Sheet position is authoritative: never process a different sheet
        # simply because its name looks like an ME2M data sheet.
        sheet = book.sheet_names[0]
        frame = pd.read_excel(book, sheet_name=sheet, dtype=object)
        names = [str(c).strip() for c in frame.columns]
        missing = [c for c in REQUIRED_COLUMNS if c not in names]
        if missing:
            raise ValueError(
                f"Required columns missing from the first worksheet '{sheet}': " + ", ".join(missing)
            )
        if len(names) != len(set(names)):
            raise ValueError("The first worksheet has duplicate column names.")
        frame.columns = names
        return frame, sheet


def _prepare_numbers(frame):
    """Convert numeric text to real values, like Excel's Convert to Number.

    Apply conversion to numeric-looking cells throughout the source sheet,
    retaining other text and leaving date columns to the date parser.
    """
    for name in frame.columns:
        if name in ("Document Date", "Delivery Date"):
            continue
        raw = frame[name]
        text = raw.astype("string").str.strip().str.removeprefix("'").str.strip()
        text = text.str.replace("\u00a0", "", regex=False).str.replace("\u202f", "", regex=False)
        text = text.str.replace(",", "", regex=False).str.replace("\u2212", "-", regex=False)
        text = text.str.replace(r"^\((.+)\)$", r"-\1", regex=True)
        text = text.str.replace(r"^(.+)-$", r"-\1", regex=True)
        values = pd.to_numeric(text.mask(text == ""), errors="coerce")
        nonfinite = values.isin([np.inf, -np.inf])
        invalid = (raw.notna() & (text != "") & values.isna()) | nonfinite
        if name in MEASURE_COLUMNS and invalid.any():
            rows = ", ".join(str(i + 2) for i in frame.index[invalid][:5])
            raise ValueError(f"Invalid number in '{name}' at Excel row(s): {rows}.")
        values = values.mask(nonfinite)
        # Text and alphanumeric identifiers remain text. A convertible value
        # replaces the cell's stored value, independently of its number format.
        frame[name] = values if name in MEASURE_COLUMNS else raw.where(values.isna(), values)


def _prepare_dates(frame):
    for name in ("Document Date", "Delivery Date"):
        if name not in frame:
            continue
        raw = frame[name]
        nonempty = raw.notna() & raw.astype(str).str.strip().ne("")
        numeric = raw.map(lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)) & nonempty
        values = pd.to_datetime(raw.where(~numeric), format="mixed", errors="coerce")
        if numeric.any():
            values.loc[numeric] = pd.to_datetime(
                raw[numeric].astype(float), unit="D", origin="1899-12-30", errors="coerce",
            )
        invalid = nonempty & values.isna()
        if invalid.any():
            rows = ", ".join(str(i + 2) for i in frame.index[invalid][:5])
            raise ValueError(f"Invalid date in '{name}' at Excel row(s): {rows}.")
        frame[name] = values


def build_report(source, output, report_date=None, allowed_codes=None, on_step=None, should_cancel=None):
    """Write Main File and Pivot Summary, and return serializable run stats."""
    as_of = report_date or report_today()
    if isinstance(as_of, str):
        as_of = date.fromisoformat(as_of)
    if isinstance(as_of, datetime):
        as_of = as_of.date()
    if not isinstance(as_of, date):
        raise ValueError("Choose a valid reporting date.")
    output = Path(output)

    def step(text):
        if should_cancel and should_cancel():
            raise ME2MCancelled()
        if on_step:
            on_step(text)

    step("Reading ME2M workbook")
    frame, sheet = _read_source(source)
    frame = frame.dropna(how="all").reset_index(drop=True)
    groups = frame["Purchasing Group"].fillna("").astype(str).str.strip().str.upper()
    source_rows = len(frame)
    if allowed_codes is not None:
        mask = groups.isin([str(c).strip().upper() for c in allowed_codes])
        frame = frame.loc[mask].copy()
        groups = groups.loc[mask]
    if frame.empty:
        raise ValueError("No ME2M rows found for your allocated purchasing groups." if allowed_codes is not None
                         else "The ME2M data sheet is empty.")

    step("Converting text numbers to Excel numbers")
    _prepare_numbers(frame)
    groups = frame["Purchasing Group"].fillna("").astype(str).str.strip().str.upper()
    step("Classifying Period using the reporting date")
    _prepare_dates(frame)
    text = frame["Delivery Text"].fillna("").astype(str)
    abc = frame["ABC Indicator"].fillna("").astype(str).str.strip().str.upper()
    cutoff = pd.Timestamp(as_of - timedelta(days=7))
    period = pd.Series(None, index=frame.index, dtype=object)
    period.loc[text.str.contains("pjw", case=False, regex=False)] = "With NCR"
    period.loc[period.isna() & text.str.contains("Product in McCook", case=False, regex=False)] = "Product in McCook"
    period.loc[period.isna() & abc.eq("X")] = "Discontinued"
    period.loc[period.isna() & frame["Delivery Date"].lt(cutoff)] = "Past Dues"
    frame["Period"] = period
    columns = [c for c in frame.columns if c != "Period"]
    columns.insert(columns.index("Delivery Date") + 1, "Period")
    frame = frame[columns]

    step("Building purchasing-group and Period summaries")
    summary = pd.DataFrame({
        "group": groups.replace("", "(blank)"), "period": period.fillna("(blank)"),
        "value": frame["Net Order Value"].fillna(0).astype(float),
    })
    pivot = summary.groupby(["group", "period"])["value"].sum().unstack(fill_value=0)
    pivot = pivot.reindex(columns=CATEGORIES, fill_value=0).sort_index()
    pivot["Grand Total"] = pivot.sum(axis=1)
    counts = {c: int(summary["period"].eq(c).sum()) for c in CATEGORIES}
    values = {c: float(pivot[c].sum()) for c in CATEGORIES}
    by_group = {}
    for group, rows in summary.groupby("group"):
        by_group[group] = {
            "rows": len(rows), "net_order_value": float(rows["value"].sum()),
            "counts": {c: int(rows["period"].eq(c).sum()) for c in CATEGORIES},
            "values": {c: float(pivot.loc[group, c]) for c in CATEGORIES},
        }
    stats = {
        "total_rows": len(frame), "source_rows": source_rows, "source_sheet": sheet,
        "net_order_value": float(summary["value"].sum()), "report_date": as_of.isoformat(),
        "cutoff_date": cutoff.date().isoformat(), "counts": counts, "values": values,
        "by_group": by_group,
    }

    step("Writing Main File and Pivot Summary")
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        with pd.ExcelWriter(output, engine="openpyxl", datetime_format="mm-dd-yy") as writer:
            frame.to_excel(writer, sheet_name="Main File", index=False)
            ws = writer.book["Main File"]
            widths = {"Short Text": 38, "Delivery Text": 58, "Vendor/supplying plant": 34}
            for cells in ws.iter_rows():
                for cell in cells:
                    # Preserve literal source text that happens to begin '='.
                    if cell.data_type == "f":
                        cell.data_type = "s"
                    if cell.row == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="305496")
                        cell.alignment = Alignment(horizontal="center", wrap_text=True)
                    elif columns[cell.column - 1] in NUMBER_FORMATS:
                        cell.number_format = NUMBER_FORMATS[columns[cell.column - 1]]
                    elif columns[cell.column - 1] in ("Document Date", "Delivery Date"):
                        cell.number_format = "mm-dd-yy"
            for i, name in enumerate(columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 22)
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            ws2 = writer.book.create_sheet("Pivot Summary")
            ws2["A1"] = f"ME2M report as of {as_of.isoformat()}"
            ws2["A2"] = f"Past Dues: Delivery Date before {stats['cutoff_date']}"
            ws2.append([])
            headers = ["Purchasing Group", *CATEGORIES, "Grand Total"]
            for j, name in enumerate(headers, 1):
                cell = ws2.cell(4, j, name)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
                ws2.column_dimensions[get_column_letter(j)].width = 24
            for i, (group, row) in enumerate(pivot.iterrows(), 5):
                ws2.cell(i, 1, group)
                for j, value in enumerate(row, 2):
                    ws2.cell(i, j, float(value)).number_format = "#,##0.00"
            last = 5 + len(pivot)
            ws2.cell(last, 1, "Grand Total")
            for j, name in enumerate(pivot.columns, 2):
                ws2.cell(last, j, float(pivot[name].sum())).number_format = "#,##0.00"
            for cell in ws2[last]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
            ws2.freeze_panes = "B5"
        step("Report ready")
    except BaseException:
        output.unlink(missing_ok=True)
        raise
    stats["output_path"] = str(output)
    return stats
