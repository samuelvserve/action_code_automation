"""ME2M rule, isolation, access and end-to-end upload regression checks."""
import importlib
import io
import json
from datetime import datetime
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from database import db
from me2m_pipeline import CATEGORIES, ME2MCancelled, build_report
from routes import me2m_routes, me2m_schedule_routes
from services import me2m_runner
from utils.auth import create_user, ensure_default_user


def workbook_bytes(rows=None, sheet="Mani File"):
    rows = rows if rows is not None else [
        {"Purchasing Group": "D13", "Delivery Date": "2026-09-30", "Delivery Text": "pjw",
         "ABC Indicator": "X", "Net Order Value": "1,200.50"},
        {"Purchasing Group": "D40", "Delivery Date": "2026-10-01", "Delivery Text": "",
         "ABC Indicator": "A", "Net Order Value": 50},
    ]
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name=sheet, index=False)
        for line in writer.book[sheet]:
            for cell in line:
                if cell.data_type == "f":
                    cell.data_type = "s"
    return stream.getvalue()


def write_source(tmp_path, rows=None, sheet="Mani File"):
    source = tmp_path / "source.xlsx"
    source.write_bytes(workbook_bytes(rows, sheet))
    return source


def test_rule_priority_cutoff_and_totals(tmp_path):
    rows = []
    for group, text, abc, day, value in [
        ("D13", "PJW Product in McCook", "X", "2026-09-30", "1,200.50"),
        ("D13", "product IN mccook", "X", "2026-09-30", 20),
        ("D13", "", "x", "2026-09-30", 30),
        ("D40", "", "A", "2026-09-30", 40),
        ("D40", "", "A", "2026-10-01", 50),
        ("D40", "", "A", "2026-10-02", 60),
        (None, "=literal source text", "A", None, 70),
    ]:
        rows.append({"Purchasing Group": group, "Delivery Text": text, "ABC Indicator": abc,
                     "Delivery Date": day, "Net Order Value": value})
    result = build_report(write_source(tmp_path, rows), tmp_path / "out.xlsx", "2026-10-08")
    assert result["counts"] == {
        "Discontinued": 1, "Past Dues": 1, "Product in McCook": 1, "With NCR": 1, "(blank)": 3,
    }
    assert result["net_order_value"] == 1470.5
    assert sum(result["values"].values()) == result["net_order_value"]
    assert sum(g["rows"] for g in result["by_group"].values()) == 7
    book = load_workbook(result["output_path"])
    assert book.sheetnames == ["Main File", "Pivot Summary"]
    sheet = book["Main File"]
    headers = [c.value for c in sheet[1]]
    assert headers.index("Period") == headers.index("Delivery Date") + 1
    val_col = headers.index("Net Order Value") + 1
    assert sheet.cell(2, val_col).value == 1200.5
    assert sheet.cell(2, val_col).number_format == "#,##0.00"
    text_col = headers.index("Delivery Text") + 1
    assert sheet.cell(8, text_col).data_type == "s"
    assert sheet.freeze_panes == "A2"
    pivot = book["Pivot Summary"]
    assert pivot.cell(pivot.max_row, 1).value == "Grand Total"
    assert pivot.cell(pivot.max_row, 7).value == 1470.5
    book.close()


def test_reporting_date_changes_only_eligible_past_dues(tmp_path):
    source = write_source(tmp_path)
    before = build_report(source, tmp_path / "before.xlsx", "2026-10-08")
    after = build_report(source, tmp_path / "after.xlsx", "2026-10-09")
    assert before["counts"]["Past Dues"] == 0
    assert after["counts"]["Past Dues"] == 1
    assert before["counts"]["With NCR"] == after["counts"]["With NCR"] == 1


def test_existing_period_replaced_and_buyer_rows_filtered(tmp_path):
    rows = pd.read_excel(io.BytesIO(workbook_bytes())).to_dict("records")
    for row in rows:
        row["Period"] = "outdated classification"
    source = write_source(tmp_path, rows, "Main File")
    result = build_report(source, tmp_path / "out.xlsx", "2026-10-08", allowed_codes=["D13"])
    assert result["total_rows"] == 1
    assert set(result["by_group"]) == {"D13"}
    with pytest.raises(ValueError, match="allocated purchasing groups"):
        build_report(source, tmp_path / "empty.xlsx", allowed_codes=["D09"])


@pytest.mark.parametrize("column,value", [("Net Order Value", "not numeric"), ("Delivery Date", "not a date")])
def test_bad_data_is_reported_not_silently_dropped(tmp_path, column, value):
    rows = pd.read_excel(io.BytesIO(workbook_bytes())).to_dict("records")
    rows[0][column] = value
    with pytest.raises(ValueError, match=column):
        build_report(write_source(tmp_path, rows), tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()


def test_invalid_sheet_and_empty_data(tmp_path):
    source = write_source(tmp_path, [{"unrelated": 1}])
    with pytest.raises(ValueError, match="Required columns"):
        build_report(source, tmp_path / "out.xlsx")
    source = write_source(tmp_path, [dict.fromkeys(["Purchasing Group", "Delivery Text", "ABC Indicator", "Delivery Date", "Net Order Value"])])
    with pytest.raises(ValueError, match="empty"):
        build_report(source, tmp_path / "out.xlsx")


def test_stop_after_write_removes_incomplete_download(tmp_path):
    progress = []
    with pytest.raises(ME2MCancelled):
        build_report(write_source(tmp_path), tmp_path / "out.xlsx", on_step=progress.append,
                     should_cancel=lambda: bool(progress) and progress[-1] == "Writing Main File and Pivot Summary")
    assert not (tmp_path / "out.xlsx").exists()


def test_first_worksheet_is_used_even_when_another_has_the_expected_name(tmp_path):
    source = tmp_path / "multiple_sheets.xlsx"
    first = pd.read_excel(io.BytesIO(workbook_bytes()), dtype=object)
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        first.to_excel(writer, sheet_name="SAP Export", index=False)
        first.assign(**{"Net Order Value": 99999}).to_excel(writer, sheet_name="Main File", index=False)
        pd.DataFrame({"Old summary": [123]}).to_excel(writer, sheet_name="Pivot Summary", index=False)
    result = build_report(source, tmp_path / "processed.xlsx", "2026-10-08")
    assert result["source_sheet"] == "SAP Export"
    assert result["total_rows"] == 2
    assert result["net_order_value"] == 1250.5
    book = load_workbook(result["output_path"])
    assert book.sheetnames == ["Main File", "Pivot Summary"]
    pivot = book.worksheets[1]
    assert pivot.cell(pivot.max_row, 7).data_type == "n"
    assert pivot.cell(pivot.max_row, 7).value == 1250.5
    book.close()


def test_first_worksheet_validation_never_falls_back_to_another_sheet(tmp_path):
    source = tmp_path / "cover_then_data.xlsx"
    with pd.ExcelWriter(source, engine="openpyxl") as writer:
        pd.DataFrame({"Cover": ["Read me"]}).to_excel(writer, sheet_name="Cover", index=False)
        pd.read_excel(io.BytesIO(workbook_bytes())).to_excel(writer, sheet_name="Main File", index=False)
    with pytest.raises(ValueError, match="first worksheet 'Cover'"):
        build_report(source, tmp_path / "processed.xlsx")
    assert not (tmp_path / "processed.xlsx").exists()


@pytest.mark.parametrize("text,expected", [
    ("'1,234.50", 1234.5), ("  '001234  ", 1234), ("1\u00a0234.50", 1234.5),
    ("1\u202f234.50", 1234.5), ("(1,234.50)", -1234.5), ("1,234.50-", -1234.5),
    ("\u22121234.50", -1234.5), ("2.5E+3", 2500),
])
def test_convert_to_number_changes_stored_excel_values(tmp_path, text, expected):
    row = {
        "Purchasing Group": "D13", "Delivery Date": "2026-09-30", "Delivery Text": "pjw",
        "ABC Indicator": "A", "Net Order Value": text,
        "Material": "'001234", "Purchasing Document": "'4501076962", "Item": "  '0020  ",
        "Plant": "1000", "Purch. Organization": "'1000", "Order Quantity": "'3,200.00",
        "Storage Location": "'0007", "Vendor Material Number": "'001234",
        "Short Text": "Bracket assembly 1234", "Material Group": "CH",
    }
    source = write_source(tmp_path, [row], "Export")
    original = load_workbook(source)
    original_headers = [c.value for c in original.worksheets[0][1]]
    assert original.worksheets[0].cell(2, original_headers.index("Net Order Value") + 1).data_type == "s"
    original.close()
    result = build_report(source, tmp_path / "processed.xlsx", "2026-10-08")
    book = load_workbook(result["output_path"])
    data = book.worksheets[0]
    cells = {header.value: cell for header, cell in zip(data[1], data[2])}
    for name, number in {
        "Net Order Value": expected, "Material": 1234, "Purchasing Document": 4501076962,
        "Item": 20, "Plant": 1000, "Purch. Organization": 1000, "Order Quantity": 3200,
        "Storage Location": 7, "Vendor Material Number": 1234,
    }.items():
        assert cells[name].data_type == "n", name
        assert cells[name].value == number, name
    assert cells["Net Order Value"].number_format == "#,##0.00"
    assert cells["Short Text"].data_type == cells["Material Group"].data_type == "s"
    assert cells["Short Text"].value == row["Short Text"]
    assert cells["Material Group"].value == "CH"
    assert cells["Delivery Date"].value.date().isoformat() == "2026-09-30"
    assert book.worksheets[1].cell(book.worksheets[1].max_row, 7).value == expected
    assert result["net_order_value"] == expected
    book.close()


@pytest.fixture
def app_env(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    for module in (me2m_routes, me2m_runner):
        for name, directory in (("ME2M_UPLOAD_DIR", "uploads"), ("ME2M_OUTPUT_DIR", "outputs"), ("ME2M_LOG_DIR", "logs")):
            monkeypatch.setattr(module, name, tmp_path / directory)
    db.init_db()
    monkeypatch.setattr(me2m_schedule_routes, "ME2M_INPUT_DIR", tmp_path / "schedule-inputs")
    ensure_default_user()
    # App startup is exercised against this isolated DB without firing any
    # automatic ACR runs or sending notifications.
    with patch("services.scheduler.start"):
        module = importlib.import_module("app")
    app = module.app
    app.config.update(TESTING=True)
    buyer_id = create_user("buyer", "test-password", "Buyer", allowed_codes=["D13"])
    other_id = create_user("other", "test-password", "Other buyer")
    admin = db.query_one("SELECT * FROM users WHERE role='admin'")
    client = app.test_client()
    with client.session_transaction() as session:
        session.update(user_id=admin["id"], username=admin["username"], role="admin", active_portal="acr")
    yield app, client, buyer_id, other_id
    for thread, _ in list(me2m_runner._RUNS.values()):
        thread.join(10)


def upload_and_wait(client, content=None, filename="ME2M.xlsx"):
    response = client.post("/me2m", data={
        "report_date": "2026-10-08", "workbook": (io.BytesIO(content if content is not None else workbook_bytes()), filename),
    }, content_type="multipart/form-data")
    assert response.status_code == 302
    run_id = int(response.location.rsplit("/", 1)[-1])
    worker = me2m_runner._RUNS.get(run_id)
    if worker:
        worker[0].join(15)
        assert not worker[0].is_alive()
    return db.query_one("SELECT * FROM me2m_runs WHERE id=?", (run_id,))


def test_upload_progress_download_stats_and_acr_isolation(app_env):
    app, client, _, _ = app_env
    db.execute("INSERT INTO runs (run_type,purchasing_codes,status,started_at,total_rows,action_items) VALUES (?,?,?,?,?,?)",
               ("Instant Run", "D13", "success", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), 777, 123))
    with patch("action_code_pipeline.run_pipeline") as acr_pipeline, patch("services.mailer.notify_run") as mailer:
        run = upload_and_wait(client)
        acr_pipeline.assert_not_called()
        mailer.assert_not_called()
    assert run["status"] == "success"
    assert run["total_rows"] == 2
    assert run["net_order_value"] == 1250.5
    assert not list(me2m_runner.ME2M_UPLOAD_DIR.iterdir())
    assert client.get(f"/me2m/{run['id']}/status").json["status"] == "success"
    assert b"Download processed XLSX" in client.get(f"/me2m/{run['id']}").data
    assert b"Period counts" in client.get(f"/me2m/{run['id']}/log").data
    download = client.get(f"/me2m/{run['id']}/download")
    assert download.status_code == 200
    assert "attachment" in download.headers["Content-Disposition"]
    download.close()
    acr_stats = client.get("/dashboard?process=acr&period=0")
    me2m_stats = client.get("/dashboard?process=me2m&period=0")
    assert acr_stats.status_code == me2m_stats.status_code == 200
    assert b"777" in acr_stats.data
    assert b"777" not in me2m_stats.data
    assert b"1,250.50" in me2m_stats.data
    filtered = client.get("/dashboard?process=me2m&period=0&group=D13")
    assert b"1,200.50" in filtered.data and b"1,250.50" not in filtered.data
    assert b"No activity in this range" in client.get("/dashboard?process=me2m&status=failed").data
    assert b"Action Code portal" in client.get("/run").data
    db.init_db()
    assert db.query_one("SELECT total_rows FROM runs")["total_rows"] == 777


def test_owner_authorization_and_buyer_workbook_scope(app_env):
    app, client, buyer_id, other_id = app_env
    admin_run = upload_and_wait(client)
    with client.session_transaction() as session:
        session.update(user_id=buyer_id, username="buyer", role="buyer")
    for suffix in ("", "/status", "/download", "/log"):
        assert client.get(f"/me2m/{admin_run['id']}{suffix}").status_code == 403
    assert client.post(f"/me2m/{admin_run['id']}/stop").status_code == 403
    assert b"ME2M.xlsx" not in client.get("/me2m").data
    own = upload_and_wait(client)
    assert own["total_rows"] == 1 and own["net_order_value"] == 1200.5
    assert set(json.loads(own["group_stats"])) == {"D13"}
    with client.session_transaction() as session:
        session.update(user_id=other_id, username="other", role="buyer")
    assert client.get(f"/me2m/{own['id']}/download").status_code == 403
    assert b"1,200.50" not in client.get("/dashboard?process=me2m&period=0").data


def test_upload_validation_failure_cleanup_and_login(app_env):
    app, client, _, _ = app_env
    for content, name, expected in ((b"abc", "file.csv", b".xlsx workbook"), (b"", "file.xlsx", b"empty")):
        response = client.post("/me2m", data={"report_date": "2026-10-08", "workbook": (io.BytesIO(content), name)})
        assert response.status_code == 200 and expected in response.data
    assert not db.query("SELECT * FROM me2m_runs")
    run = upload_and_wait(client, b"not an Excel workbook")
    assert run["status"] == "failed" and run["error"]
    assert not list(me2m_runner.ME2M_UPLOAD_DIR.iterdir())
    assert client.get(f"/me2m/{run['id']}/download").status_code == 404
    with client.session_transaction() as session:
        session.clear()
    assert client.get("/me2m").status_code == 302
    assert client.get("/dashboard?process=me2m").status_code == 302


def test_download_path_guard(app_env):
    _, client, _, _ = app_env
    run = upload_and_wait(client)
    db.execute("UPDATE me2m_runs SET output_file='../test.db', log_file='../test.db' WHERE id=?", (run["id"],))
    assert client.get(f"/me2m/{run['id']}/download").status_code == 404
    assert client.get(f"/me2m/{run['id']}/log").status_code == 404


def test_blank_purchasing_group_dashboard_filter(app_env):
    _, client, _, _ = app_env
    row = {"Purchasing Group": None, "Delivery Date": "2026-10-02", "Delivery Text": "",
           "ABC Indicator": "A", "Net Order Value": 99}
    upload_and_wait(client, workbook_bytes([row]))
    response = client.get("/dashboard?process=me2m&period=0&group=(blank)")
    assert response.status_code == 200
    assert b"99.00" in response.data
    assert b"99.00" in response.data


def test_upload_size_limit(app_env, monkeypatch):
    _, client, _, _ = app_env
    monkeypatch.setattr(me2m_routes, "ME2M_MAX_UPLOAD_BYTES", 10)
    response = client.post("/me2m", data={"report_date": "2026-10-08", "workbook": (io.BytesIO(workbook_bytes()), "test.xlsx")})
    assert response.status_code == 200
    assert b"50 MB or smaller" in response.data
    assert not db.query("SELECT * FROM me2m_runs")
    assert not list(me2m_runner.ME2M_UPLOAD_DIR.iterdir())
